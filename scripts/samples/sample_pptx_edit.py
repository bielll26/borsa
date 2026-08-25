#!/usr/bin/env python3
"""
Sample Script: Edit PowerPoint (PPTX) Files
Usage: python scripts/samples/sample_pptx_edit.py
Requires: pip install python-pptx
RULE: NEVER modify original files. Always copy first.
"""
import os
import shutil
from datetime import date

PROJECT_ROOT = r"C:\Users\Anit\Downloads\10x-Content-Expert"
INPUT_DIR = os.path.join(PROJECT_ROOT, "input")
WORKING_DIR = os.path.join(PROJECT_ROOT, "output", "working")
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output", "pptx")

def safe_copy(source_path):
    """Copy file to working directory. NEVER edit originals."""
    os.makedirs(WORKING_DIR, exist_ok=True)
    filename = os.path.basename(source_path)
    name, ext = os.path.splitext(filename)
    copy_path = os.path.join(WORKING_DIR, f"{name}_copy{ext}")
    shutil.copy2(source_path, copy_path)
    print(f"Safe copy created: {copy_path}")
    return copy_path

def analyze_presentation(pptx_path):
    """Analyze PPTX structure: slides, text, layout."""
    try:
        from pptx import Presentation
    except ImportError:
        print("ERROR: Install python-pptx first: pip install python-pptx")
        return []

    prs = Presentation(pptx_path)
    slides_info = []

    for i, slide in enumerate(prs.slides, 1):
        slide_data = {"number": i, "texts": [], "layout": slide.slide_layout.name}
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_data["texts"].append(text)
        slides_info.append(slide_data)
        print(f"Slide {i} ({slide_data['layout']}): {len(slide_data['texts'])} text elements")

    print(f"\nTotal: {len(slides_info)} slides")
    return slides_info

def replace_text_in_slide(pptx_path, slide_number, old_text, new_text, output_name):
    """Replace text on a specific slide. Preserves formatting."""
    try:
        from pptx import Presentation
    except ImportError:
        print("ERROR: Install python-pptx first: pip install python-pptx")
        return ""

    prs = Presentation(pptx_path)
    slide = prs.slides[slide_number - 1]  # 1-indexed input
    replaced = False

    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if old_text in run.text:
                        run.text = run.text.replace(old_text, new_text)
                        replaced = True

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    today = date.today().strftime("%Y-%m-%d")
    output_path = os.path.join(OUTPUT_DIR, f"{today}_{output_name}_modified.pptx")
    prs.save(output_path)

    if replaced:
        print(f"Replaced '{old_text}' with '{new_text}' on slide {slide_number}")
    else:
        print(f"WARNING: '{old_text}' not found on slide {slide_number}")
    print(f"Saved to: {output_path}")
    return output_path

def extract_all_text(pptx_path):
    """Extract all text from all slides."""
    try:
        from pptx import Presentation
    except ImportError:
        print("ERROR: Install python-pptx first: pip install python-pptx")
        return ""

    prs = Presentation(pptx_path)
    all_text = ""

    for i, slide in enumerate(prs.slides, 1):
        all_text += f"\n=== Slide {i} ===\n"
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text += shape.text_frame.text + "\n"

    return all_text

if __name__ == "__main__":
    source = os.path.join(INPUT_DIR, "example.pptx")

    if os.path.exists(source):
        working_copy = safe_copy(source)
        slides = analyze_presentation(working_copy)
        text = extract_all_text(working_copy)
        print(f"\nAll text:\n{text[:500]}...")
    else:
        print(f"No file at {source}")
        print("Place a PPTX in input/ folder and update the 'source' variable")
