#!/usr/bin/env python3
"""
Sample Script: Edit Word (DOCX) Files
Usage: python scripts/samples/sample_docx_edit.py
Requires: pip install python-docx
RULE: NEVER modify original files. Always copy first.
"""
import os
import shutil
from datetime import date

PROJECT_ROOT = r"C:\Users\Anit\Downloads\10x-Content-Expert"
INPUT_DIR = os.path.join(PROJECT_ROOT, "input")
WORKING_DIR = os.path.join(PROJECT_ROOT, "output", "working")
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output", "docx")

def safe_copy(source_path):
    """Copy file to working directory. NEVER edit originals."""
    os.makedirs(WORKING_DIR, exist_ok=True)
    filename = os.path.basename(source_path)
    name, ext = os.path.splitext(filename)
    copy_path = os.path.join(WORKING_DIR, f"{name}_copy{ext}")
    shutil.copy2(source_path, copy_path)
    print(f"Safe copy created: {copy_path}")
    return copy_path

def analyze_document(docx_path):
    """Analyze DOCX structure: headings, paragraphs, tables."""
    try:
        from docx import Document
    except ImportError:
        print("ERROR: Install python-docx first: pip install python-docx")
        return {}

    doc = Document(docx_path)
    info = {"paragraphs": 0, "headings": [], "tables": len(doc.tables), "words": 0}

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        info["paragraphs"] += 1
        info["words"] += len(text.split())
        if para.style.name.startswith("Heading"):
            info["headings"].append({"level": para.style.name, "text": text})

    print(f"Document: {info['paragraphs']} paragraphs, {info['words']} words, {info['tables']} tables")
    for h in info["headings"]:
        print(f"  {h['level']}: {h['text']}")
    return info

def find_and_replace(docx_path, old_text, new_text, output_name):
    """Find and replace text in DOCX. Preserves formatting."""
    try:
        from docx import Document
    except ImportError:
        print("ERROR: Install python-docx first: pip install python-docx")
        return ""

    doc = Document(docx_path)
    count = 0

    for para in doc.paragraphs:
        for run in para.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)
                count += 1

    # Also check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if old_text in run.text:
                            run.text = run.text.replace(old_text, new_text)
                            count += 1

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    today = date.today().strftime("%Y-%m-%d")
    output_path = os.path.join(OUTPUT_DIR, f"{today}_{output_name}_modified.docx")
    doc.save(output_path)

    print(f"Replaced {count} occurrence(s) of '{old_text}'")
    print(f"Saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    source = os.path.join(INPUT_DIR, "example.docx")

    if os.path.exists(source):
        working_copy = safe_copy(source)
        info = analyze_document(working_copy)
    else:
        print(f"No file at {source}")
        print("Place a DOCX in input/ folder and update the 'source' variable")
