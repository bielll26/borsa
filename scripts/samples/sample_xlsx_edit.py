#!/usr/bin/env python3
"""
Sample Script: Edit Excel (XLSX) Files
Usage: python scripts/samples/sample_xlsx_edit.py
Requires: pip install openpyxl
RULE: NEVER modify original files. Always copy first.
"""
import os
import shutil
from datetime import date

PROJECT_ROOT = r"C:\Users\Anit\Downloads\10x-Content-Expert"
INPUT_DIR = os.path.join(PROJECT_ROOT, "input")
WORKING_DIR = os.path.join(PROJECT_ROOT, "output", "working")
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output", "xlsx")

def safe_copy(source_path):
    """Copy file to working directory. NEVER edit originals."""
    os.makedirs(WORKING_DIR, exist_ok=True)
    filename = os.path.basename(source_path)
    name, ext = os.path.splitext(filename)
    copy_path = os.path.join(WORKING_DIR, f"{name}_copy{ext}")
    shutil.copy2(source_path, copy_path)
    print(f"Safe copy created: {copy_path}")
    return copy_path

def analyze_spreadsheet(xlsx_path):
    """Analyze XLSX: sheets, dimensions, data."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: Install openpyxl first: pip install openpyxl")
        return {}

    wb = load_workbook(xlsx_path, read_only=True)
    info = {"sheets": []}

    for name in wb.sheetnames:
        ws = wb[name]
        sheet_info = {
            "name": name,
            "rows": ws.max_row or 0,
            "cols": ws.max_column or 0,
        }
        info["sheets"].append(sheet_info)
        print(f"Sheet '{name}': {sheet_info['rows']} rows x {sheet_info['cols']} cols")

    wb.close()
    return info

def update_cell(xlsx_path, sheet_name, cell_ref, new_value, output_name):
    """Update a specific cell value."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: Install openpyxl first: pip install openpyxl")
        return ""

    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name]
    old_value = ws[cell_ref].value
    ws[cell_ref] = new_value

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    today = date.today().strftime("%Y-%m-%d")
    output_path = os.path.join(OUTPUT_DIR, f"{today}_{output_name}_modified.xlsx")
    wb.save(output_path)
    wb.close()

    print(f"Updated {cell_ref}: '{old_value}' -> '{new_value}'")
    print(f"Saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    source = os.path.join(INPUT_DIR, "example.xlsx")

    if os.path.exists(source):
        working_copy = safe_copy(source)
        info = analyze_spreadsheet(working_copy)
    else:
        print(f"No file at {source}")
        print("Place an XLSX in input/ folder and update the 'source' variable")
