#!/usr/bin/env python3
"""
Create Vibe Marketing Presentation
Generates a 5-slide high-energy PPTX about Vibe Marketing possibilities
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

def create_vibe_marketing_presentation(output_path: str):
    """Create the Vibe Marketing presentation"""

    # Create presentation with widescreen dimensions
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Get blank layout
    blank_layout = prs.slide_layouts[6]  # Blank layout

    # Colors
    DARK_BG = RGBColor(18, 18, 18)
    ELECTRIC_BLUE = RGBColor(0, 200, 255)
    NEON_PINK = RGBColor(255, 0, 128)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(200, 200, 200)
    GOLD = RGBColor(255, 215, 0)

    def add_background(slide, color=DARK_BG):
        """Add dark background shape"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0,
            prs.slide_width, prs.slide_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        # Send to back
        spTree = slide.shapes._spTree
        sp = shape._element
        spTree.remove(sp)
        spTree.insert(2, sp)

    def add_text_box(slide, text, left, top, width, height,
                     font_size=24, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
        """Add formatted text box"""
        txBox = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
        return txBox

    # ============================================
    # SLIDE 1: THE HOOK
    # ============================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    # Main title
    add_text_box(slide1, "Marketing Just Changed Forever.",
                 0.5, 1.5, 12, 1.2, font_size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Subtitle
    add_text_box(slide1, "The old playbook is dead.",
                 0.5, 2.8, 12, 0.6, font_size=28, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    # Vibe Marketing intro
    add_text_box(slide1, "Welcome to the era of VIBE MARKETING.",
                 0.5, 3.6, 12, 0.8, font_size=36, bold=True, color=ELECTRIC_BLUE, align=PP_ALIGN.CENTER)

    # Three pillars
    pillars = "Where AI meets creativity.\nWhere speed meets strategy.\nWhere one person can outperform entire teams."
    add_text_box(slide1, pillars,
                 0.5, 4.6, 12, 1.5, font_size=24, color=WHITE, align=PP_ALIGN.CENTER)

    # CTA
    add_text_box(slide1, "Are you ready?",
                 0.5, 6.3, 12, 0.6, font_size=32, bold=True, color=NEON_PINK, align=PP_ALIGN.CENTER)

    # ============================================
    # SLIDE 2: THE SHIFT
    # ============================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)

    # Title
    add_text_box(slide2, "From Months to Minutes",
                 0.5, 0.5, 12, 1, font_size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # OLD WAY - Left side
    add_text_box(slide2, "OLD WAY",
                 0.8, 1.6, 5, 0.5, font_size=28, bold=True, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    old_way = "• Weeks to plan a campaign\n• Days to write copy\n• Hours to create variations\n• Endless approval cycles"
    add_text_box(slide2, old_way,
                 0.8, 2.2, 5.5, 2.5, font_size=22, color=LIGHT_GRAY, align=PP_ALIGN.LEFT)

    # VIBE MARKETING WAY - Right side
    add_text_box(slide2, "VIBE MARKETING",
                 7, 1.6, 5.5, 0.5, font_size=28, bold=True, color=ELECTRIC_BLUE, align=PP_ALIGN.CENTER)

    new_way = "• Idea to execution in HOURS\n• 100 variations in minutes\n• Test everything, keep what works\n• Move at the speed of thought"
    add_text_box(slide2, new_way,
                 7, 2.2, 5.5, 2.5, font_size=22, color=WHITE, align=PP_ALIGN.LEFT)

    # Bottom statement
    add_text_box(slide2, "This isn't automation. This is AMPLIFICATION.",
                 0.5, 5.8, 12, 0.8, font_size=32, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

    # ============================================
    # SLIDE 3: WHAT'S POSSIBLE NOW
    # ============================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)

    # Title
    add_text_box(slide3, "The New Marketing Superpowers",
                 0.5, 0.4, 12, 0.8, font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Superpowers list
    powers = [
        ("1. Launch campaigns 10X faster", "From concept to live in a single day"),
        ("2. Create personalized content at scale", "Speak to 1,000 audiences like they're the only one"),
        ("3. Test ideas without fear", "Generate dozens of hooks, headlines, and angles instantly"),
        ("4. Repurpose everything", "One video becomes 50 pieces of content"),
        ("5. Stay consistently creative", "AI handles execution, you drive vision"),
    ]

    y_pos = 1.3
    for title, desc in powers:
        add_text_box(slide3, title, 0.8, y_pos, 11, 0.5, font_size=24, bold=True, color=ELECTRIC_BLUE)
        add_text_box(slide3, desc, 0.8, y_pos + 0.4, 11, 0.4, font_size=18, color=LIGHT_GRAY)
        y_pos += 1.0

    # Bottom statement
    add_text_box(slide3, "Your competition is still in meetings. You're already live.",
                 0.5, 6.5, 12, 0.6, font_size=26, bold=True, color=NEON_PINK, align=PP_ALIGN.CENTER)

    # ============================================
    # SLIDE 4: THE ENERGY (V.I.B.E.)
    # ============================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)

    # Title
    add_text_box(slide4, "This Is Marketing On Fire",
                 0.5, 0.4, 12, 0.8, font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # VIBE acronym
    vibe_items = [
        ("V", "Velocity", "move fast, learn faster"),
        ("I", "Intelligence", "AI-powered insights"),
        ("B", "Bold", "take creative risks"),
        ("E", "Execution", "ideas mean nothing without action"),
    ]

    y_pos = 1.5
    for letter, word, desc in vibe_items:
        add_text_box(slide4, letter, 1.5, y_pos, 0.6, 0.6, font_size=36, bold=True, color=NEON_PINK)
        add_text_box(slide4, f"- {word}", 2.2, y_pos, 3, 0.6, font_size=28, bold=True, color=WHITE)
        add_text_box(slide4, f"({desc})", 5.5, y_pos + 0.05, 5, 0.5, font_size=20, color=LIGHT_GRAY)
        y_pos += 0.75

    # Formula
    add_text_box(slide4, "When you combine:",
                 0.5, 4.5, 12, 0.5, font_size=22, color=WHITE, align=PP_ALIGN.CENTER)

    formula = "Human creativity + AI capability\nStrategic thinking + Instant execution\nBrand vision + Unlimited content"
    add_text_box(slide4, formula,
                 0.5, 5.0, 12, 1.2, font_size=20, color=ELECTRIC_BLUE, align=PP_ALIGN.CENTER)

    # Bottom statement
    add_text_box(slide4, "You don't just do marketing. You become unstoppable.",
                 0.5, 6.4, 12, 0.7, font_size=28, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

    # ============================================
    # SLIDE 5: YOUR MOVE
    # ============================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)

    # Title
    add_text_box(slide5, "The Future Belongs to the Fast",
                 0.5, 0.4, 12, 0.8, font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    add_text_box(slide5, "Right now, you have two choices:",
                 0.5, 1.2, 12, 0.5, font_size=24, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    # Choice 1 - Left
    add_text_box(slide5, "1. Watch from the sidelines",
                 0.8, 2.0, 5.5, 0.5, font_size=24, bold=True, color=LIGHT_GRAY)
    choice1 = "• Wait for the \"perfect\" moment\n• Let others figure it out first\n• Fall behind while studying the playbook"
    add_text_box(slide5, choice1,
                 0.8, 2.6, 5.5, 1.8, font_size=18, color=LIGHT_GRAY)

    # Choice 2 - Right
    add_text_box(slide5, "2. Jump in and VIBE",
                 7, 2.0, 5.5, 0.5, font_size=24, bold=True, color=ELECTRIC_BLUE)
    choice2 = "• Start experimenting TODAY\n• Use AI to 10X your output\n• Build while others are still planning"
    add_text_box(slide5, choice2,
                 7, 2.6, 5.5, 1.8, font_size=18, color=WHITE)

    # Bottom statements
    add_text_box(slide5, "The masterclass is over.",
                 0.5, 5.5, 12, 0.6, font_size=28, color=WHITE, align=PP_ALIGN.CENTER)

    add_text_box(slide5, "The real work starts NOW.",
                 0.5, 6.2, 12, 0.7, font_size=36, bold=True, color=NEON_PINK, align=PP_ALIGN.CENTER)

    # Save
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))

    return str(output)


if __name__ == '__main__':
    output_file = "output/pptx/vibe-marketing-possibilities.pptx"
    result = create_vibe_marketing_presentation(output_file)
    print(f"Presentation created: {result}")
